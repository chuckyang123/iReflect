"""
@changelog
| Version | Description                                      | Reference                                   |
| v1.0.0  | Tests: permissions, sheets, history, limits      | REQ: 20260929-Excel导出功能升级 TECH: 04_design_tech-design.md §3.4 |
/@changelog

@author chuckyang123
"""
import io
import json
from unittest import mock

from django.test import TransactionTestCase, override_settings
from openpyxl import load_workbook

from courses.models import (
    Course, CourseGroup, CourseMembership, CourseMilestone, CourseSubmission,
    Role, SubmissionType,
)
from users.models import AccountType, User
from .milestone_excel_export import (
    CELL_TEXT_LIMIT,
    TRUNCATION_SUFFIX,
    TARGET_QUESTION_TEXT,
    AUTO_GEN_QUESTION_SHEET,
)
from .models import AIFeedbackRecord, FeedbackAnswerVersion, FeedbackInitialResponse
from .research_export import export_guard
from .test_ai_feedback_verification import AIFeedbackBase, jwt_client

URL = "/api/feedback/milestone-export/"
SUBMISSIONS = "Submissions"
HISTORY = "AI Feedback History"


@override_settings(EXCEL_EXPORT_MAX_SECONDS=60)
class MilestoneExcelExportTests(TransactionTestCase):
    setUp = AIFeedbackBase.setUp

    def record(self, answer, number, question="Why?", genre=None, mechanic=None, **kwargs):
        """Create one answer version plus the feedback record attached to it."""
        version = FeedbackAnswerVersion.objects.create(
            course=self.course, milestone=self.milestone, template=self.template,
            submission=self.submission, creator=self.student_membership,
            name="Reflection", question=question, answer_content=answer,
            version_number=number, genre=genre, mechanic=mechanic,
        )
        fields = {
            "version": version, "feedback_type": AIFeedbackRecord.FEEDBACK_TYPE_REFLECTION,
            "strategy": AIFeedbackRecord.STRATEGY_BASIC, "feedback_content": "Explain more.",
            "score_json": {"stage_1": 1}, "model_version": "test-model",
            "token_usage_json": {"prompt_tokens": 7}, "latency_ms": 10,
        }
        fields.update(kwargs)
        return AIFeedbackRecord.objects.create(**fields)

    def download(self, user=None, **params):
        payload = {"course_id": self.course.id, "milestone_id": self.milestone.id}
        payload.update(params)
        return jwt_client(user or self.educator).post(URL, payload, format="json")

    def sheets(self, response):
        """Return {sheet title: [{column: value}]} and release every resource."""
        self.assertEqual(response.status_code, 200, getattr(response, "data", None))
        try:
            body = b"".join(response.streaming_content)
        finally:
            response.close()
        workbook = load_workbook(io.BytesIO(body), read_only=True)
        try:
            return {
                title: self.as_dicts(workbook[title]) for title in workbook.sheetnames
            }
        finally:
            workbook.close()

    def as_dicts(self, sheet):
        rows = list(sheet.iter_rows(values_only=True))
        header = list(rows[0]) if rows else []
        # openpyxl trims trailing empty cells, so a row whose last columns are
        # blank is shorter than the header and zip() silently drops those
        # columns. Pad to the header length so every column is always present.
        return [
            dict(zip(header, list(row) + [None] * (len(header) - len(row))))
            for row in rows[1:]
        ]

    def test_teacher_downloads_workbook_with_submissions_and_headers(self):
        self.form.form_field_data = [
            {"label": "Why?", "type": "TEXTAREA"},
            {"label": "Plan", "type": "TEXT"},
        ]
        self.form.save()
        self.submission.form_response_data = [
            {"label": "Why?", "response": "第一版答案"},
            {"label": "Plan", "response": ["A", "B"]},
        ]
        self.submission.save()

        response = self.download()
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertEqual(response["Cache-Control"], "no-store, private")
        self.assertEqual(response["X-Content-Type-Options"], "nosniff")
        self.assertTrue(
            response["Content-Disposition"].startswith(
                f'attachment; filename="ireflect-{self.course.id}-{self.milestone.id}-'
            )
        )
        self.assertTrue(response["Content-Disposition"].endswith('.xlsx"'))

        data = self.sheets(response)
        self.assertEqual(list(data), [SUBMISSIONS, HISTORY])
        self.assertEqual(data[HISTORY], [])
        row = data[SUBMISSIONS][0]
        self.assertEqual(row["Student ID"], self.student.id)
        self.assertEqual(row["Student Name"], "Student")
        self.assertEqual(row["Student Email"], "student@test.com")
        self.assertIsNone(row["Group"])
        self.assertEqual(row["Submission ID"], self.submission.id)
        self.assertEqual(row["Submission Name"], "My Submission")
        self.assertEqual(row["Submission Type"], SubmissionType.INDIVIDUAL)
        self.assertEqual(row["Status"], "Submitted")
        self.assertEqual(row["Last Editor"], "Student")
        self.assertEqual(row["Why?"], "第一版答案")
        self.assertEqual(row["Plan"], "A; B")

    def test_auto_generated_question_sheet_includes_answer(self):
        # PRD 20260901: the reflective question is not a form field. Its answer
        # is stored on the field that owns it, as "finalizationReason".
        self.form.form_field_data = [
            {"label": "Why?", "type": "TEXTAREA", "hasFeedback": True},
            {"label": "Plan", "type": "TEXT"},
        ]
        self.form.save()
        self.submission.form_response_data = [
            {
                "label": "Why?",
                "type": "TEXTAREA",
                "response": "第一版答案",
                "finalizationReason": "Because the workflow reached ARCHIVE_COMPLETE.",
            },
            {"label": "Plan", "type": "TEXT", "response": ["A", "B"]},
        ]
        self.submission.save()
        self.record("v1 answer", 1, question="Why?")

        data = self.sheets(self.download())
        self.assertIn(AUTO_GEN_QUESTION_SHEET, data)
        rows = data[AUTO_GEN_QUESTION_SHEET]
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row["Student ID"], self.student.id)
        self.assertEqual(row["Field"], "Why?")
        self.assertEqual(row["Question"], TARGET_QUESTION_TEXT)
        self.assertEqual(row["Student Answer"], "Because the workflow reached ARCHIVE_COMPLETE.")
        self.assertEqual(row["AI Feedback"], "Explain more.")

    def test_auto_generated_question_found_via_feedback_response(self):
        # The reflective question is auto-generated and need not exist as a
        # static form field; only the reflection flow recorded it. The sheet
        # must still appear with the student answer.
        self.form.form_field_data = [{"label": "Plan", "type": "TEXT"}]
        self.form.save()
        FeedbackInitialResponse.objects.create(
            course=self.course, milestone=self.milestone, template=self.template,
            creator=self.student_membership, name="Reflection",
            question=TARGET_QUESTION_TEXT,
            initial_response="The paragraph showed the game's induced emotions clearly.",
        )

        data = self.sheets(self.download())
        self.assertIn(AUTO_GEN_QUESTION_SHEET, data)
        row = data[AUTO_GEN_QUESTION_SHEET][0]
        self.assertEqual(row["Question"], TARGET_QUESTION_TEXT)
        self.assertEqual(
            row["Student Answer"],
            "The paragraph showed the game's induced emotions clearly.",
        )

    def test_auto_generated_question_accepts_underscoreized_key(self):
        # The DRF camel-case parser underscoreizes nested JSON on save, so the
        # same answer can come back stored as finalization_reason.
        self.form.form_field_data = [
            {"label": "Why?", "type": "TEXT_AREA", "hasFeedback": True},
        ]
        self.form.save()
        self.submission.form_response_data = [
            {
                "label": "Why?",
                "type": "TEXT_AREA",
                "response": "ans",
                "finalization_reason": "Saved with the underscoreized key.",
            },
        ]
        self.submission.save()

        data = self.sheets(self.download())
        self.assertIn(AUTO_GEN_QUESTION_SHEET, data)
        row = data[AUTO_GEN_QUESTION_SHEET][0]
        self.assertEqual(row["Field"], "Why?")
        self.assertEqual(
            row["Student Answer"], "Saved with the underscoreized key."
        )

    def test_draft_submission_marked(self):
        self.submission.is_draft = True
        self.submission.save()
        self.assertEqual(self.sheets(self.download())[SUBMISSIONS][0]["Status"], "Draft")

    def test_students_and_unrelated_educators_denied(self):
        self.assertEqual(self.download(self.student).status_code, 403)
        outsider = User.objects.create(email="outside@example.com", account_type=AccountType.EDUCATOR, is_activated=True)
        self.assertEqual(self.download(outsider).status_code, 403)

    def test_standard_instructor_and_admin_can_export(self):
        instructor = User.objects.create(email="invited@example.com", account_type=AccountType.STANDARD, is_activated=True)
        CourseMembership.objects.create(course=self.course, user=instructor, role=Role.INSTRUCTOR)
        self.sheets(self.download(instructor))
        self.educator.account_type = AccountType.ADMIN
        self.educator.save()
        self.sheets(self.download())

    def test_milestone_must_belong_to_the_requested_course(self):
        other_course = Course.objects.create(owner=self.educator, name="Other", description="", is_published=True)
        other_milestone = CourseMilestone.objects.create(
            course=other_course, name="M2", description="", is_published=True,
            start_date_time=self.milestone.start_date_time,
            end_date_time=self.milestone.end_date_time,
        )
        self.assertEqual(self.download(milestone_id=other_milestone.id).status_code, 403)
        self.assertEqual(self.download(course_id=other_course.id).status_code, 403)

    def test_group_submission_uses_group_identity(self):
        group = CourseGroup.objects.create(course=self.course, name="Team A")
        CourseSubmission.objects.create(
            course=self.course, milestone=self.milestone, template=self.template,
            creator=self.student_membership, editor=self.student_membership, group=group,
            name="Group Work", description="", is_draft=False,
            submission_type=SubmissionType.GROUP, form_response_data=[],
        )
        rows = self.sheets(self.download())[SUBMISSIONS]
        row = next(item for item in rows if item["Submission Name"] == "Group Work")
        self.assertEqual(row["Student ID"], self.student.id)
        self.assertEqual(row["Student Name"], "Team A")
        self.assertEqual(row["Student Email"], "Team A")
        self.assertEqual(row["Group"], "Team A")
        self.assertEqual(row["Submission Type"], SubmissionType.GROUP)

    def test_feedback_history_row_carries_version_answer_and_score(self):
        record = self.record("v1 answer", 1)
        FeedbackInitialResponse.objects.create(
            course=self.course, milestone=self.milestone, template=self.template,
            creator=self.student_membership, name="Reflection", question="Why?",
            initial_response="initial text",
        )
        row = self.sheets(self.download())[HISTORY][0]
        self.assertEqual(row["Record ID"], record.id)
        self.assertEqual(row["Student ID"], self.student.id)
        self.assertEqual(row["Student Name"], "Student")
        self.assertEqual(row["Submission ID"], self.submission.id)
        self.assertEqual(row["Question"], "Why?")
        self.assertEqual(row["Version Number"], 1)
        self.assertEqual(row["Answer At Version"], "v1 answer")
        self.assertEqual(row["Initial Response"], "initial text")
        self.assertEqual(row["Feedback Type"], AIFeedbackRecord.FEEDBACK_TYPE_REFLECTION)
        self.assertEqual(row["Model"], "test-model")
        self.assertEqual(row["Prompt Tokens"], 7)
        self.assertIsNone(row["Completion Tokens"])
        self.assertIsNone(row["Total Tokens"])
        self.assertEqual(row["Latency (ms)"], 10)
        self.assertEqual(json.loads(row["Score JSON"]), {"stage_1": 1})

    def test_feedback_history_row_carries_genre_and_mechanic(self):
        # Genre/Main Mechanic are chosen by the student and drive the
        # knowledge-graph half of the playtest score, so a reviewer needs them
        # next to the grading. They live on the answer version, not the record.
        self.record("v1 answer", 1, genre="Action", mechanic="Mobility & Movement")
        row = self.sheets(self.download())[HISTORY][0]
        self.assertEqual(row["Genre"], "Action")
        self.assertEqual(row["Main Mechanic"], "Mobility & Movement")

    def test_feedback_history_genre_and_mechanic_blank_when_not_playtested(self):
        # Reflection feedback never goes through playtest, so both columns must
        # stay empty instead of breaking the export.
        self.record("v1 answer", 1)
        row = self.sheets(self.download())[HISTORY][0]
        self.assertIsNone(row["Genre"])
        self.assertIsNone(row["Main Mechanic"])

    def test_token_usage_keys_tolerated(self):
        self.record("camel", 1, token_usage_json={"promptTokens": 5, "completionTokens": 2, "totalTokens": 7})
        self.record("empty", 2, token_usage_json=None)
        rows = self.sheets(self.download())[HISTORY]
        self.assertEqual((rows[0]["Prompt Tokens"], rows[0]["Total Tokens"]), (5, 7))
        self.assertEqual(rows[0]["Answer At Version"], "camel")
        self.assertIsNone(rows[1]["Prompt Tokens"])
        self.assertIsNone(rows[1]["Total Tokens"])

    def test_feedback_rows_sorted_by_question_then_version(self):
        self.record("b1", 1, question="B")
        self.record("a1", 1, question="A")
        self.record("a2", 2, question="A")
        rows = self.sheets(self.download())[HISTORY]
        self.assertEqual(
            [(row["Question"], row["Version Number"], row["Answer At Version"]) for row in rows],
            [("A", 1, "a1"), ("A", 2, "a2"), ("B", 1, "b1")],
        )

    def test_long_answer_truncated_with_marker(self):
        self.form.form_field_data = [{"label": "Why?", "type": "TEXTAREA"}]
        self.form.save()
        self.submission.form_response_data = [
            {"label": "Why?", "response": "x" * (CELL_TEXT_LIMIT + 100)}
        ]
        self.submission.save()
        cell = self.sheets(self.download())[SUBMISSIONS][0]["Why?"]
        self.assertEqual(len(cell), CELL_TEXT_LIMIT)
        self.assertTrue(cell.endswith(TRUNCATION_SUFFIX))

    @override_settings(EXCEL_EXPORT_MAX_ROWS=0)
    def test_limit_returns_error_not_partial_workbook(self):
        response = self.download()
        self.assertEqual(response.status_code, 413)
        self.assertFalse(response.streaming)

    def test_duplicate_export_does_not_build_again(self):
        with export_guard(self.educator.id), mock.patch("feedback.milestone_excel_export.build_workbook") as build:
            self.assertEqual(self.download().status_code, 409)
            build.assert_not_called()

    def test_camel_case_request_from_frontend(self):
        client = jwt_client(self.educator)
        response = client.post(
            URL, {"courseId": self.course.id, "milestoneId": self.milestone.id}, format="json"
        )
        self.sheets(response)

    def test_get_cannot_start_export(self):
        self.assertEqual(jwt_client(self.educator).get(URL).status_code, 405)
